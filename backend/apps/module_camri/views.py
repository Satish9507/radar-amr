import io
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from django.http import HttpResponse

from .calculator import parse_camri_excel, validate_sites, run_camri
from .coefficients import ARB_COEFFICIENTS, ARG_COEFFICIENTS
from apps.config.models import ModuleConfig

_ACCEPTED = ('.xlsx', '.xls')


def _check_file(file, label):
    if not file:
        return f'No {label} file provided. Send as multipart/form-data.'
    if not any(file.name.endswith(ext) for ext in _ACCEPTED):
        return f'{label} file must be .xlsx or .xls.'
    return None


# ── Validate ───────────────────────────────────────────────────────────────────

class CAMRIValidateView(APIView):
    """
    Validates one or both CAMRI files.
    - Single file: checks columns against known coefficients.
    - Both files: also checks Sample_ID matching between ARB and ARG.
    """
    def post(self, request):
        arb_file = request.FILES.get('arb_file')
        arg_file = request.FILES.get('arg_file')

        if not arb_file and not arg_file:
            return Response(
                {'status': 'error', 'code': 'NO_FILE',
                 'message': 'Provide at least one file: arb_file or arg_file.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        for f, label in [(arb_file, 'ARB'), (arg_file, 'ARG')]:
            if f is not None:
                err = _check_file(f, label)
                if err:
                    return Response(
                        {'status': 'error', 'code': 'NO_FILE', 'message': err},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

        arb_records = arb_matched = arb_unmatched = None
        arg_records = arg_matched = arg_unmatched = None

        try:
            if arb_file:
                arb_records, arb_matched, arb_unmatched, _ = parse_camri_excel(arb_file, 'arb')
            if arg_file:
                arg_records, arg_matched, arg_unmatched, _ = parse_camri_excel(arg_file, 'arg')
        except ValueError as exc:
            return Response(
                {'status': 'error', 'code': 'PARSE_ERROR', 'message': str(exc)},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except Exception as exc:
            return Response(
                {'status': 'error', 'code': 'PARSE_ERROR',
                 'message': f'Failed to read file: {exc}'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        site_report = None
        if arb_records and arg_records:
            site_report = validate_sites(arb_records, arg_records)

        return Response({
            'status':             'success',
            'mode':               'combined' if (arb_file and arg_file) else ('arb_only' if arb_file else 'arg_only'),
            'site_validation':    site_report,
            'arb_matched_cols':   arb_matched,
            'arb_unmatched_cols': arb_unmatched,
            'arg_matched_cols':   arg_matched,
            'arg_unmatched_cols': arg_unmatched,
            'arb_sample_count':   len(arb_records) if arb_records else None,
            'arg_sample_count':   len(arg_records) if arg_records else None,
        })


# ── Calculate ──────────────────────────────────────────────────────────────────

class CAMRICalculateView(APIView):
    def post(self, request):
        arb_file = request.FILES.get('arb_file')
        arg_file = request.FILES.get('arg_file')

        if not arb_file and not arg_file:
            return Response(
                {'status': 'error', 'code': 'NO_FILE',
                 'message': 'Provide at least one file: arb_file or arg_file.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        for f, label in [(arb_file, 'ARB'), (arg_file, 'ARG')]:
            if f is not None:
                err = _check_file(f, label)
                if err:
                    return Response(
                        {'status': 'error', 'code': 'NO_FILE', 'message': err},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

        try:
            alpha = float(request.data.get('alpha', 0.6))
            alpha = max(0.5, min(1.0, alpha))
        except (TypeError, ValueError):
            alpha = 0.6

        arb_records = arb_matched = arb_unmatched = arb_warns = None
        arg_records = arg_matched = arg_unmatched = arg_warns = None

        try:
            if arb_file:
                arb_records, arb_matched, arb_unmatched, arb_warns = parse_camri_excel(arb_file, 'arb')
            if arg_file:
                arg_records, arg_matched, arg_unmatched, arg_warns = parse_camri_excel(arg_file, 'arg')
        except ValueError as exc:
            return Response(
                {'status': 'error', 'code': 'PARSE_ERROR', 'message': str(exc)},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except Exception as exc:
            return Response(
                {'status': 'error', 'code': 'PARSE_ERROR',
                 'message': f'Failed to read file: {exc}'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Validate matching only when both files are provided
        site_report = None
        if arb_records and arg_records:
            site_report = validate_sites(arb_records, arg_records)
            if not site_report['valid']:
                return Response(
                    {'status': 'error', 'code': 'NO_MATCHED_SITES',
                     'message': 'No matching Sample_IDs found between ARB and ARG files.',
                     'site_validation': site_report},
                    status=status.HTTP_400_BAD_REQUEST,
                )

        try:
            try:
                cfg = ModuleConfig.objects.get(module_code='CAMRI')
                thresholds = cfg.risk_thresholds.get('levels') if cfg.risk_thresholds else None
            except ModuleConfig.DoesNotExist:
                thresholds = None

            payload = run_camri(
                arb_records=arb_records or [],
                arg_records=arg_records or [],
                alpha=alpha,
                thresholds=thresholds,
            )
        except Exception as exc:
            return Response(
                {'status': 'error', 'code': 'CALC_ERROR', 'message': str(exc)},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        warnings = list(arb_warns or []) + list(arg_warns or [])
        if arb_unmatched:
            warnings.append({
                'code':    'UNMATCHED_ARB_COLS',
                'message': f'ARB columns with no burden coefficient (excluded): '
                           f'{", ".join(arb_unmatched)}',
            })
        if arg_unmatched:
            warnings.append({
                'code':    'UNMATCHED_ARG_COLS',
                'message': f'ARG columns with no burden coefficient (excluded): '
                           f'{", ".join(arg_unmatched)}',
            })

        return Response({
            'status':  'success',
            'module':  'CAMRI',
            'version': '1.0',
            'meta': {
                'mode':           payload['mode'],
                'k':              payload['k'],
                'alpha':          payload.get('alpha'),
                'processed_at':   datetime.now(timezone.utc).isoformat(),
                'arb_matched':    arb_matched,
                'arg_matched':    arg_matched,
                'site_validation': site_report,
            },
            **payload,
            'warnings': warnings,
        })


# ── Excel templates ────────────────────────────────────────────────────────────

_HEADER_FILL   = PatternFill('solid', fgColor='7C3AED')
_HEADER_FONT   = Font(color='FFFFFF', bold=True, size=10)
_UNIT_FILL     = PatternFill('solid', fgColor='EDE9FE')
_UNIT_FONT     = Font(color='5B21B6', italic=True, size=9)
_META_COLS     = ['Sample_ID', 'Site', 'Month', 'Date', 'Category', 'Sub_category']


def _build_template(data_cols: list, unit_label: str, example_rows: list) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active

    all_cols = _META_COLS + data_cols

    # Row 1 — units
    for col_idx, col in enumerate(all_cols, start=1):
        cell = ws.cell(row=1, column=col_idx)
        if col not in _META_COLS:
            cell.value = unit_label
        cell.fill = _UNIT_FILL
        cell.font = _UNIT_FONT
        cell.alignment = Alignment(horizontal='center')

    # Row 2 — headers
    for col_idx, col in enumerate(all_cols, start=1):
        cell = ws.cell(row=2, column=col_idx, value=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal='center')

    # Rows 3+ — example data
    for row_idx, row_data in enumerate(example_rows, start=3):
        for col_idx, col in enumerate(all_cols, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(col, ''))

    # Column widths
    for col_idx, col in enumerate(all_cols, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = (
            18 if col in _META_COLS else 12
        )

    ws.freeze_panes = 'G3'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class CAMRITemplateView(APIView):
    def get(self, request, file_type):
        if file_type == 'arb':
            data_cols   = list(ARB_COEFFICIENTS.keys())
            unit_label  = 'CFU/mL'
            filename    = 'camri_arb_template.xlsx'
            example_rows = [
                {'Sample_ID': 'SG-001', 'Site': 'Kallang River', 'Month': 'Dec',
                 'Date': '2015-12-15', 'Category': 'Tributary', 'Sub_category': 'Urban',
                 'Ec_CAZ': 89.5, 'Pseu_MEM': 548.5, 'Kleb_CAZ': 39.7,
                 'Kleb_MEM': 27.1, 'Ent_VAN': 29.6, 'Ec_MEM': 110.6},
                {'Sample_ID': 'SG-002', 'Site': 'Bedok Reservoir', 'Month': 'Dec',
                 'Date': '2015-12-15', 'Category': 'Freshwater Lake', 'Sub_category': 'Urban',
                 'Ec_CAZ': 0.58, 'Pseu_MEM': 24.2, 'Kleb_CAZ': 0.36,
                 'Kleb_MEM': 5.13, 'Ent_VAN': 0.04, 'Ec_MEM': 13.8},
                {'Sample_ID': 'SG-003', 'Site': 'MacRitchie Res', 'Month': 'Dec',
                 'Date': '2015-12-15', 'Category': 'Freshwater Lake', 'Sub_category': 'Parkland',
                 'Ec_CAZ': 0.0, 'Pseu_MEM': 0.18, 'Kleb_CAZ': 0.09,
                 'Kleb_MEM': 0.0, 'Ent_VAN': 0.0, 'Ec_MEM': 0.0},
            ]
        elif file_type == 'arg':
            data_cols   = list(ARG_COEFFICIENTS.keys())
            unit_label  = 'copies/mL'
            filename    = 'camri_arg_template.xlsx'
            example_rows = [
                {'Sample_ID': 'SG-001', 'Site': 'Kallang River', 'Month': 'Dec',
                 'Date': '2015-12-15', 'Category': 'Tributary', 'Sub_category': 'Urban',
                 'blaKPC': 60.0, 'blaCTX_M': 0.77, 'vanA': 10.3, 'blaNDM': 10.6,
                 'blaSHV': 40.3, 'tetO': 1829.1, 'tetM': 50.4, 'qnrA': 71.2},
                {'Sample_ID': 'SG-002', 'Site': 'Bedok Reservoir', 'Month': 'Dec',
                 'Date': '2015-12-15', 'Category': 'Freshwater Lake', 'Sub_category': 'Urban',
                 'blaKPC': 54.6, 'blaCTX_M': 0.04, 'vanA': 2.53, 'blaNDM': 8.8,
                 'blaSHV': 6.06, 'tetO': 6.29, 'tetM': 0.0, 'qnrA': 110.2},
                {'Sample_ID': 'SG-003', 'Site': 'MacRitchie Res', 'Month': 'Dec',
                 'Date': '2015-12-15', 'Category': 'Freshwater Lake', 'Sub_category': 'Parkland',
                 'blaKPC': 12.0, 'blaCTX_M': 0.0, 'vanA': 0.5, 'blaNDM': 2.1,
                 'blaSHV': 1.2, 'tetO': 1.0, 'tetM': 0.0, 'qnrA': 8.3},
            ]
        else:
            return Response(
                {'status': 'error', 'message': 'file_type must be "arb" or "arg"'},
                status=status.HTTP_404_NOT_FOUND,
            )

        xlsx_bytes = _build_template(data_cols, unit_label, example_rows)
        response = HttpResponse(
            xlsx_bytes,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
